# -*- coding: utf-8 -*-
import sys
from pptx import Presentation

sys.stdout.reconfigure(encoding="utf-8")
prs = Presentation("presentation/Prezentatsia_Reshetov_Sentiment_Analysis.pptx")
print("slides", len(prs.slides), "size", prs.slide_width, "x", prs.slide_height)
for i, s in enumerate(prs.slides):
    titles = []
    for sh in s.shapes:
        if sh.has_text_frame and sh.text_frame.text.strip():
            titles.append(sh.text_frame.text.strip().split("\n")[0][:65])
    print(f"{i+1}. {titles[0] if titles else '-'}")
