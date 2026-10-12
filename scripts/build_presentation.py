# -*- coding: utf-8 -*-
"""Презентация: стиль Степанова (16:9, панели, логотип) + академичность Акчулпановой."""
from __future__ import annotations

import json
import shutil
from pathlib import Path

import pandas as pd
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Emu, Pt

ROOT = Path(__file__).resolve().parents[1]
MEDIA = ROOT / "presentation" / "_stepanov_media"
FIGURES = ROOT / "figures"
RESULTS = ROOT / "results"
OUT = ROOT / "presentation" / "Prezentatsia_Reshetov_Sentiment_Analysis.pptx"

SW, SH = 9144000, 5143500

summary = json.loads((RESULTS / "data_summary.json").read_text(encoding="utf-8"))
metrics = pd.read_csv(RESULTS / "model_metrics.csv")
calibration = json.loads((RESULTS / "calibration_metrics.json").read_text(encoding="utf-8"))
report = pd.read_csv(RESULTS / "classification_report.csv", index_col=0)
best = metrics.iloc[0]


def fmt(x, nd=3):
    return f"{float(x):.{nd}f}".replace(".", ",")


def set_run(run, size=18, bold=False, color=(30, 30, 30)):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = "Calibri"
    run.font.color.rgb = RGBColor(*color)


def add_text(slide, left, top, width, height, text, size=18, bold=False,
             color=(30, 30, 30), align=PP_ALIGN.LEFT):
    box = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(str(text).split("\n")):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_after = Pt(4)
        run = p.add_run()
        run.text = line
        set_run(run, size=size, bold=bold, color=color)
    return box


def add_bullets(slide, left, top, width, height, items, size=15):
    box = slide.shapes.add_textbox(Emu(left), Emu(top), Emu(width), Emu(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = PP_ALIGN.LEFT
        p.space_after = Pt(10)
        run = p.add_run()
        run.text = "•  " + item
        set_run(run, size=size, color=(40, 40, 40))
    return box


def add_logo(slide):
    logo = MEDIA / "image1.png"
    if logo.exists():
        slide.shapes.add_picture(str(logo), Emu(6457575), Emu(4791026), width=Emu(2686425))


def add_left_panel(slide):
    panel = MEDIA / "image4.png"
    if panel.exists():
        slide.shapes.add_picture(str(panel), Emu(0), Emu(0), height=Emu(SH))


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def add_table(slide, left, top, width, height, data):
    rows, cols = len(data), len(data[0])
    table = slide.shapes.add_table(rows, cols, Emu(left), Emu(top), Emu(width), Emu(height)).table
    for i, row in enumerate(data):
        for j, val in enumerate(row):
            cell = table.cell(i, j)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                p.alignment = PP_ALIGN.CENTER
                for run in p.runs:
                    set_run(run, size=12, bold=(i == 0), color=(25, 25, 25))
    return table


def build():
    prs = Presentation()
    prs.slide_width = SW
    prs.slide_height = SH

    # 1. Титульный
    s = blank(prs)
    side = MEDIA / "image14.png"
    if side.exists():
        s.shapes.add_picture(str(side), Emu(5395608), Emu(0), height=Emu(SH))
    add_text(
        s, 350000, 200000, 4900000, 1500000,
        "Министерство науки и высшего образования\nРоссийской Федерации\n"
        "ФГБОУ ВО «Оренбургский государственный\nуниверситет имени В. А. Бондаренко»\n"
        "Институт математики и информационных технологий\nКафедра прикладной математики",
        size=11, color=(70, 70, 70), align=PP_ALIGN.CENTER,
    )
    add_text(s, 350000, 1750000, 4900000, 400000, "КУРСОВАЯ РАБОТА",
             size=22, bold=True, color=(20, 40, 80), align=PP_ALIGN.CENTER)
    add_text(s, 350000, 2150000, 4900000, 300000, "по дисциплине «Машинное обучение»",
             size=14, color=(55, 55, 55), align=PP_ALIGN.CENTER)
    add_text(
        s, 300000, 2550000, 5000000, 1200000,
        "Применение методов машинного обучения\nдля автоматической классификации\nи анализа тональности отзывов клиентов",
        size=17, bold=True, color=(20, 20, 20), align=PP_ALIGN.CENTER,
    )
    add_text(
        s, 350000, 3950000, 4900000, 900000,
        "Автор: студент группы 25ИСТ(м)ИИП — Решетов Е. В.\n"
        "Руководитель: канд. техн. наук, доцент С. Т. Дусакаева\nОренбург, 2026",
        size=12, color=(50, 50, 50), align=PP_ALIGN.CENTER,
    )

    # 2. Актуальность
    s = blank(prs)
    add_left_panel(s)
    add_logo(s)
    add_text(s, 2800000, 350000, 5800000, 500000, "Актуальность и цель исследования",
             size=24, bold=True, color=(20, 40, 80))
    add_bullets(
        s, 2800000, 1000000, 5800000, 3500000,
        [
            "Поток клиентских отзывов слишком велик для единообразной ручной разметки.",
            "Автоматическая тональность нужна для маршрутизации негативных обращений.",
            "Русский язык требует учёта морфологии, отрицаний и возможного шума меток.",
            "Цель — сравнительное исследование методов ML для трёхклассовой классификации тональности на корпусе RuReviews.",
            "Задачи: обзор методов; постановка задачи; эксперимент со сравнением моделей; анализ результатов и ограничений.",
        ],
        size=14,
    )

    # 3. Данные
    s = blank(prs)
    add_logo(s)
    add_text(s, 450000, 180000, 8200000, 400000, "Данные: корпус RuReviews",
             size=24, bold=True, color=(20, 40, 80))
    add_bullets(
        s, 450000, 600000, 8200000, 1300000,
        [
            f"Русскоязычные отзывы о товарах (женская одежда и аксессуары); исходных записей: {summary['initial_rows']}.",
            f"После очистки: {summary['modeling_rows']} (коротких: {summary['short_texts_removed']}, дубликатов: {summary['duplicate_texts_removed']}).",
            f"Классы: негатив {summary['class_counts']['negative']}, нейтрально {summary['class_counts']['neutral']}, позитив {summary['class_counts']['positive']}.",
            f"Разбиение 60/20/20: {summary['train_rows']} / {summary['validation_rows']} / {summary['test_rows']}; в модель — только текст.",
        ],
        size=13,
    )
    fig = FIGURES / "class_distribution.png"
    if fig.exists():
        s.shapes.add_picture(str(fig), Emu(1100000), Emu(2100000), width=Emu(6900000))
    add_text(s, 1100000, 4650000, 6900000, 280000,
             "Рисунок 1 – Распределение классов после очистки",
             size=11, color=(90, 90, 90), align=PP_ALIGN.CENTER)

    # 4. Методы
    s = blank(prs)
    add_logo(s)
    add_text(s, 450000, 300000, 8200000, 450000, "Методы и протокол эксперимента",
             size=24, bold=True, color=(20, 40, 80))
    add_bullets(
        s, 450000, 900000, 8200000, 3600000,
        [
            "Векторизация: TF-IDF (словные n-граммы 1–2; для LinearSVC — символьные 3–5).",
            "Модели: Dummy, MultinomialNB, LogisticRegression, LinearSVC, Random Forest, MLP.",
            "Подбор гиперпараметров только по validation macro-F1; словарь TF-IDF — внутри Pipeline.",
            "Финальное обучение на train+validation; итоговая оценка — один раз на test.",
            "Дополнительно: матрица ошибок, топ-признаки, сигмоидальная калибровка вероятностей.",
        ],
        size=15,
    )

    # 5. Метрики
    s = blank(prs)
    add_left_panel(s)
    add_logo(s)
    add_text(s, 2850000, 350000, 5700000, 500000, "Метрики качества",
             size=24, bold=True, color=(20, 40, 80))
    add_bullets(
        s, 2850000, 1000000, 5700000, 3400000,
        [
            "Accuracy — доля всех верных ответов.",
            "Macro-F1 — основная метрика (равный вес трёх классов).",
            "Precision / Recall — анализ ошибок по классам.",
            "ROC-AUC (OvR, macro) — качество ранжирования вероятностей.",
            "Log loss / Brier score — качество вероятностных оценок.",
            "F1 = 2 · (Precision · Recall) / (Precision + Recall).",
        ],
        size=15,
    )

    # 6. Результаты
    s = blank(prs)
    add_logo(s)
    add_text(s, 450000, 220000, 8200000, 400000, "Результаты сравнения моделей",
             size=24, bold=True, color=(20, 40, 80))
    table_data = [["Модель", "Accuracy", "Macro-F1", "ROC-AUC", "Log loss"]]
    for _, row in metrics.iterrows():
        table_data.append([
            str(row["model"]),
            fmt(row["accuracy"]),
            fmt(row["f1_macro"]),
            fmt(row["roc_auc_ovr_macro"]),
            fmt(row["log_loss"], 3),
        ])
    add_table(s, 450000, 800000, 8200000, 2800000, table_data)
    add_text(
        s, 450000, 3750000, 8200000, 800000,
        f"Таблица 1 – Метрики на отложенном тесте.\n"
        f"Лучшая модель — LogisticRegression: accuracy = {fmt(best.accuracy)}, "
        f"macro-F1 = {fmt(best.f1_macro)}, ROC-AUC = {fmt(best.roc_auc_ovr_macro)}.",
        size=13, color=(35, 35, 35),
    )

    # 7. Матрица ошибок
    s = blank(prs)
    add_logo(s)
    add_text(s, 400000, 200000, 5000000, 450000, "Матрица ошибок лучшей модели",
             size=22, bold=True, color=(20, 40, 80))
    add_bullets(
        s, 400000, 750000, 4700000, 3200000,
        [
            "Правильно: негатив 4 066, нейтрально 3 887, позитив 4 976.",
            "Основная ошибка — смешение с нейтральным классом.",
            "Путаница крайних классов (негатив ↔ позитив) невелика.",
            f"Позитив: F1 ≈ {fmt(report.loc['Позитив', 'f1-score'])}; нейтраль: F1 ≈ {fmt(report.loc['Нейтрально', 'f1-score'])}.",
            "Часть высокоуверенных ошибок связана с шумом меток рейтинга.",
        ],
        size=14,
    )
    cm = FIGURES / "confusion_matrix.png"
    if cm.exists():
        s.shapes.add_picture(str(cm), Emu(5200000), Emu(700000), width=Emu(3600000))
    add_text(s, 5200000, 4300000, 3600000, 300000,
             "Рисунок 2 – Матрица ошибок LogisticRegression",
             size=11, color=(90, 90, 90), align=PP_ALIGN.CENTER)

    # 8. Интерпретация
    s = blank(prs)
    add_logo(s)
    add_text(s, 400000, 180000, 8200000, 400000, "Интерпретация признаков",
             size=24, bold=True, color=(20, 40, 80))
    add_bullets(
        s, 400000, 650000, 8200000, 1100000,
        [
            "Коэффициенты логистической регрессии согласуются с тональной лексикой.",
            "Негатив — отрицания и сильная негативная оценка; позитив — благодарность и высокая оценка.",
            "Интерпретируемость линейной модели — практическое преимущество перед «чёрным ящиком».",
        ],
        size=14,
    )
    top = FIGURES / "top_features.png"
    if top.exists():
        s.shapes.add_picture(str(top), Emu(600000), Emu(1850000), width=Emu(7900000))
    add_text(s, 600000, 4550000, 7900000, 280000,
             "Рисунок 3 – Наиболее значимые n-граммы LogisticRegression",
             size=11, color=(90, 90, 90), align=PP_ALIGN.CENTER)

    # 9. Калибровка
    s = blank(prs)
    add_logo(s)
    add_text(s, 400000, 180000, 8200000, 400000, "Калибровка вероятностей",
             size=24, bold=True, color=(20, 40, 80))
    add_bullets(
        s, 400000, 650000, 8200000, 1200000,
        [
            f"Brier до калибровки: {fmt(calibration['multiclass_brier_before'], 4)}; после sigmoid CV: {fmt(calibration['multiclass_brier_after'], 4)}.",
            f"Log loss: {fmt(calibration['log_loss_before'], 4)} → {fmt(calibration['log_loss_after'], 4)}.",
            "Посткалибровка ухудшила метрики и отклонена; исходные вероятности LogReg уже близки к диагонали.",
            "На практике: пороги уверенности + ручная проверка спорных случаев.",
        ],
        size=14,
    )
    cal = FIGURES / "calibration_curve.png"
    if cal.exists():
        s.shapes.add_picture(str(cal), Emu(700000), Emu(2000000), width=Emu(7700000))
    add_text(s, 700000, 4550000, 7700000, 280000,
             "Рисунок 4 – Калибровочные кривые до и после постобработки",
             size=11, color=(90, 90, 90), align=PP_ALIGN.CENTER)

    # 10. Заключение
    s = blank(prs)
    add_left_panel(s)
    add_logo(s)
    right = MEDIA / "image16.png"
    if right.exists():
        s.shapes.add_picture(str(right), Emu(7200000), Emu(900000), height=Emu(3200000))
    add_text(s, 2800000, 300000, 4300000, 450000, "Заключение и практическая значимость",
             size=22, bold=True, color=(20, 40, 80))
    add_bullets(
        s, 2800000, 900000, 4300000, 3600000,
        [
            f"Лучшая модель — LogisticRegression: macro-F1 = {fmt(best.f1_macro)}, ROC-AUC = {fmt(best.roc_auc_ovr_macro)}.",
            "Линейная модель превзошла Random Forest и MLP на выбранном протоколе.",
            "Наиболее сложен нейтральный класс; сигмоидальная калибровка не потребовалась.",
            "Применение: маршрутизация негативных отзывов в поддержку интернет-магазина.",
            "Далее: внешняя проверка на другом домене и сравнение с RuBERT.",
        ],
        size=13,
    )

    # 11. Источники
    s = blank(prs)
    add_logo(s)
    add_text(s, 450000, 300000, 8200000, 450000, "Список использованных источников",
             size=24, bold=True, color=(20, 40, 80))
    add_bullets(
        s, 450000, 900000, 8200000, 3600000,
        [
            "Smetanin S., Komarov M. Sentiment Analysis of Product Reviews in Russian… IEEE CBI, 2019.",
            "Salton G., Buckley C. Term-weighting approaches in automatic text retrieval. 1988.",
            "Cortes C., Vapnik V. Support-vector networks. Machine Learning, 1995.",
            "Devlin J. et al. BERT… NAACL-HLT, 2019.",
            "Pedregosa F. et al. Scikit-learn: Machine Learning in Python. JMLR, 2011.",
            "Breiman L. Random Forests. Machine Learning, 2001; Hastie et al. ESL, 2009.",
        ],
        size=13,
    )

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    shutil.copy2(OUT, ROOT / "Prezentatsia_Reshetov_Sentiment_Analysis.pptx")
    print("Saved", OUT, "slides", len(prs.slides))


if __name__ == "__main__":
    build()
