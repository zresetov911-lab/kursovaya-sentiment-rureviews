# -*- coding: utf-8 -*-
import sys
from pathlib import Path
from docx import Document

sys.stdout.reconfigure(encoding="utf-8")
doc = Document(Path(__file__).resolve().parents[1] / "report" / "Kursovaya_Reshetov_Sentiment_Analysis.docx")
print("total paras", len(doc.paragraphs))
for i, p in enumerate(doc.paragraphs):
    t = p.text
    if "риложение" in t.lower() or t.strip().startswith("from __future__") or t.strip().startswith("import json"):
        fonts = [(r.font.name, r.font.size.pt if r.font.size else None) for r in p.runs[:1]]
        print(i, fonts, repr(t[:120]))

# show block around last 'Приложение'
idxs = [i for i, p in enumerate(doc.paragraphs) if "Приложение А" in p.text or "ПРИЛОЖЕНИЕ А" in p.text]
print("appendix idxs", idxs)
if idxs:
    start = idxs[-1]
    for j in range(start, min(start + 20, len(doc.paragraphs))):
        p = doc.paragraphs[j]
        fonts = [(r.font.name, r.font.size.pt if r.font.size else None) for r in p.runs[:1]]
        print(j, fonts, repr(p.text[:100]))
